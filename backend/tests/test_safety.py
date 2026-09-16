"""Guardrails added for deployment (docs/decisions.md D-041): rate limits, secret
redaction, security headers and CORS, and upload hardening."""

from __future__ import annotations

import io
import zipfile

import pytest

from app.core.config import get_settings
from app.core.ratelimit import SlidingWindowLimiter
from app.rag.parse import UnsupportedDocument, check_file

FAKE_GOOGLE_KEY = "AIza" + "B" * 35


# --- 1. Rate limits ---------------------------------------------------------


class Clock:
    def __init__(self):
        self.now = 1000.0

    def __call__(self):
        return self.now


def test_the_minute_limit_refuses_the_next_request_and_says_how_long_to_wait():
    clock = Clock()
    limiter = SlidingWindowLimiter(clock=clock)
    limits = [(2, 60.0), (100, 3600.0)]
    assert limiter.hit("ask", "u1", limits) is None
    clock.now += 10
    assert limiter.hit("ask", "u1", limits) is None
    clock.now += 5
    assert limiter.hit("ask", "u1", limits) == pytest.approx(45.0)
    clock.now += 45
    assert limiter.hit("ask", "u1", limits) is None


def test_the_hour_limit_applies_even_when_the_minute_limit_would_allow():
    clock = Clock()
    limiter = SlidingWindowLimiter(clock=clock)
    limits = [(5, 60.0), (3, 3600.0)]
    for _ in range(3):
        assert limiter.hit("chat", "u1", limits) is None
        clock.now += 120
    assert limiter.hit("chat", "u1", limits) is not None


def test_a_refused_request_does_not_extend_the_wait():
    clock = Clock()
    limiter = SlidingWindowLimiter(clock=clock)
    limits = [(1, 60.0), (100, 3600.0)]
    limiter.hit("ask", "u1", limits)
    for _ in range(5):
        clock.now += 1
        limiter.hit("ask", "u1", limits)
    clock.now = 1060.0
    assert limiter.hit("ask", "u1", limits) is None


def test_limits_are_per_user_and_per_bucket():
    limiter = SlidingWindowLimiter(clock=Clock())
    limits = [(1, 60.0), (100, 3600.0)]
    assert limiter.hit("ask", "u1", limits) is None
    assert limiter.hit("ask", "u2", limits) is None
    assert limiter.hit("chat", "u1", limits) is None
    assert limiter.hit("ask", "u1", limits) is not None


def _chat_setup(monkeypatch, decisions, db_tables=None):
    from app.agent import loop as loop_module
    from app.api import agent_chat
    from app.core.supabase import service_db, user_db
    from app.core.workspace import get_workspace_context
    from app.main import app
    from tests.test_agent import TODAY, FakeDb, ScriptedGateway, make_ctx

    db, service = FakeDb(db_tables or {"tasks": []}), FakeDb()
    app.dependency_overrides[get_workspace_context] = lambda: make_ctx("x").workspace

    async def _db():
        yield db

    async def _service():
        yield service

    app.dependency_overrides[user_db] = _db
    app.dependency_overrides[service_db] = _service
    monkeypatch.setattr(agent_chat, "today", lambda: TODAY)
    gateway = ScriptedGateway(decisions)
    monkeypatch.setattr(loop_module, "get_gateway", lambda: gateway)
    return gateway, service


def test_the_assistant_endpoint_returns_429_with_retry_after(client, monkeypatch):
    from app.core import ratelimit

    monkeypatch.setattr(ratelimit, "_limits", lambda bucket: [(1, 60.0), (100, 3600.0)])
    _chat_setup(monkeypatch, [{"action": "respond", "response": "Hi."}] * 2)
    assert client.post("/agent/chat", json={"message": "hello"}).status_code == 200
    second = client.post("/agent/chat", json={"message": "hello again"})
    assert second.status_code == 429
    assert int(second.headers["retry-after"]) >= 1
    assert "Try again in" in second.json()["detail"]


# --- 2. Secret redaction ----------------------------------------------------


def test_a_pasted_key_never_reaches_the_model_or_the_audit_log(client, monkeypatch):
    gateway, service = _chat_setup(monkeypatch, [{"action": "respond", "response": "Noted."}])
    response = client.post("/agent/chat", json={
        "message": f"My key is {FAKE_GOOGLE_KEY}, what do I have to do?",
        "history": [{"role": "user", "content": f"earlier I pasted {FAKE_GOOGLE_KEY}"}],
    })
    assert response.status_code == 200
    sent_to_model = " ".join(" ".join(call["parts"]) for call in gateway.calls)
    assert FAKE_GOOGLE_KEY not in sent_to_model
    assert "[redacted Google API key]" in sent_to_model
    audit_row = next(row for table, row in service.inserts if table == "audit_log")
    assert FAKE_GOOGLE_KEY not in str(audit_row)
    assert audit_row["details"]["secrets_redacted"] == ["Google API key"]


def test_a_key_in_the_assistants_reply_is_redacted(client, monkeypatch):
    _chat_setup(monkeypatch, [{"action": "respond", "response": f"The document lists {FAKE_GOOGLE_KEY}."}])
    body = client.post("/agent/chat", json={"message": "what does the config doc say?"}).json()
    assert FAKE_GOOGLE_KEY not in body["reply"]


# --- 3. Security headers and CORS -------------------------------------------


def test_api_responses_carry_security_headers(client):
    headers = client.get("/health").headers
    assert headers["x-content-type-options"] == "nosniff"
    assert headers["x-frame-options"] == "DENY"
    assert headers["content-security-policy"] == "default-src 'none'; frame-ancestors 'none'"
    assert headers["cache-control"] == "no-store"
    assert headers["referrer-policy"] == "no-referrer"


def test_preflight_responses_carry_them_too_and_foreign_origins_are_refused(client):
    ok = client.options("/health", headers={"Origin": "http://localhost:5173", "Access-Control-Request-Method": "GET"})
    assert ok.headers["access-control-allow-origin"] == "http://localhost:5173"
    assert ok.headers["x-content-type-options"] == "nosniff"
    foreign = client.options("/health", headers={"Origin": "https://evil.example", "Access-Control-Request-Method": "GET"})
    assert "access-control-allow-origin" not in foreign.headers


def test_several_origins_are_accepted_but_a_wildcard_is_refused(monkeypatch):
    settings = get_settings()
    monkeypatch.setattr(settings, "frontend_origin", "https://meridian.app, https://preview.meridian.app/")
    assert settings.cors_origins == ["https://meridian.app", "https://preview.meridian.app"]
    monkeypatch.setattr(settings, "frontend_origin", "*")
    with pytest.raises(ValueError):
        _ = settings.cors_origins


# --- 4. Upload hardening ----------------------------------------------------


def _pdf(pages: int) -> bytes:
    import pymupdf

    doc = pymupdf.open()
    for i in range(pages):
        doc.new_page().insert_text((72, 72), f"Page {i + 1}")
    return doc.tobytes()


def _zip(entries: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    return buffer.getvalue()


def test_a_real_pdf_passes():
    check_file("pdf", _pdf(2))


def test_a_file_named_pdf_that_is_not_a_pdf_is_refused():
    with pytest.raises(UnsupportedDocument, match="not a PDF"):
        check_file("pdf", b"<html><script>alert(1)</script></html>")


def test_a_pdf_over_the_page_limit_is_refused(monkeypatch):
    monkeypatch.setattr(get_settings(), "max_pdf_pages", 1)
    with pytest.raises(UnsupportedDocument, match="2 pages"):
        check_file("pdf", _pdf(2))


def test_a_file_named_docx_that_is_not_a_zip_is_refused():
    with pytest.raises(UnsupportedDocument, match="not one"):
        check_file("docx", b"%PDF-1.7 pretending")


def test_a_zip_that_is_not_a_word_document_is_refused():
    with pytest.raises(UnsupportedDocument, match="not a Word document"):
        check_file("docx", _zip({"payload.txt": b"hello"}))


def test_a_docx_that_expands_past_the_limit_is_refused(monkeypatch):
    monkeypatch.setattr(get_settings(), "max_docx_uncompressed_bytes", 1024 * 1024)
    bomb = _zip({"word/document.xml": b"0" * (2 * 1024 * 1024)})  # compresses to a few KB
    assert len(bomb) < 50_000
    with pytest.raises(UnsupportedDocument, match="expands to 2 MB"):
        check_file("docx", bomb)


def test_a_docx_with_too_many_parts_is_refused(monkeypatch):
    monkeypatch.setattr(get_settings(), "max_docx_entries", 3)
    with pytest.raises(UnsupportedDocument, match="too many"):
        check_file("docx", _zip({"word/document.xml": b"x", **{f"word/media/{i}.png": b"x" for i in range(5)}}))


def test_a_document_over_the_passage_limit_fails_ingestion_with_a_clear_reason(monkeypatch):
    import docx

    from app.rag.ingest import _parse_and_chunk

    document = docx.Document()
    for i in range(40):
        document.add_heading(f"Section {i}", level=1)
        document.add_paragraph("Long passage text about maintenance intervals. " * 40)
    buffer = io.BytesIO()
    document.save(buffer)
    monkeypatch.setattr(get_settings(), "max_passages_per_document", 3)
    with pytest.raises(UnsupportedDocument, match="the limit is 3"):
        _parse_and_chunk("docx", buffer.getvalue(), "big.docx")


def test_the_upload_endpoint_refuses_a_disguised_file_before_touching_storage_or_the_database(client, as_admin):
    from app.core.security import AuthenticatedUser, get_current_user
    from app.main import app

    # as_admin already makes any database call fail the test; storage is never reached either.
    app.dependency_overrides[get_current_user] = lambda: AuthenticatedUser(
        id="00000000-0000-0000-0000-000000000002", email="person@example.com", token="t")
    response = client.post(
        "/documents/upload",
        files={"file": ("report.pdf", b"MZ\x90\x00 not a pdf", "application/pdf")},
    )
    assert response.status_code == 415
    assert "not a PDF" in response.json()["detail"]
