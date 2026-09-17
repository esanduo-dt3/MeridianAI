"""Parsing and chunking. The central property: every chunk's text is exactly its
span of the canonical document text, so citations highlight the right passage."""

from io import BytesIO

import docx
import pymupdf
import pytest

from app.rag.chunker import chunk_document, estimate_tokens
from app.rag.documents import ParsedDocument, code, heading, para, table
from app.rag.parse import UnsupportedDocument, detect_type, parse_file

LONG_SENTENCE_PROSE = " ".join(
    f"Rectifier bank {i} at the Galle substation moves to a six month inspection interval after the firmware upgrade."
    for i in range(1, 90)
)
NO_PUNCTUATION = " ".join(f"token{i}" for i in range(4000))


def assert_offsets(text: str, chunks) -> None:
    assert chunks, "expected at least one chunk"
    for chunk in chunks:
        assert text[chunk.start : chunk.end] == chunk.text, f"chunk {chunk.index} text does not match its span"
        assert chunk.text == chunk.text.strip(), "chunk spans must not include surrounding whitespace"
        assert 0 <= chunk.start < chunk.end <= len(text)
    assert [c.index for c in chunks] == list(range(len(chunks)))


def make_docx() -> bytes:
    document = docx.Document()
    document.add_heading("Maintenance Handbook", 0)
    document.add_heading("Rectifiers", 1)
    document.add_paragraph(LONG_SENTENCE_PROSE)
    rows = [["Site", "Units", "Interval"]] + [[f"Site {i}", str(i * 2), "6 months"] for i in range(1, 160)]
    grid = document.add_table(rows=len(rows), cols=3)
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            grid.cell(r, c).text = value
    document.add_heading("Batteries", 1)
    document.add_paragraph("Replace batteries every four years. Keep the replacement log in the site binder.")
    document.add_paragraph("Short.")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_pdf() -> bytes:
    doc = pymupdf.open()
    for page_no in range(1, 4):
        page = doc.new_page()
        page.insert_text((72, 60), f"Operations Manual - page {page_no}", fontsize=9)
        page.insert_text((72, 110), f"Section {page_no} Overview", fontsize=18)
        y = 150
        for line in range(14):
            page.insert_text(
                (72, y), f"Line {line} of section {page_no} explains the rectifier maintenance procedure in detail.", fontsize=10
            )
            y += 16
    data = doc.tobytes()
    doc.close()
    return data


def test_detect_type_accepts_pdf_and_docx_only():
    assert detect_type("a.PDF", None) == "pdf"
    assert detect_type("report", "application/pdf") == "pdf"
    assert detect_type("notes.docx", None) == "docx"
    with pytest.raises(UnsupportedDocument):
        detect_type("sheet.xlsx", "application/vnd.ms-excel")


def test_docx_offsets_sections_and_table_splitting():
    parsed = parse_file("docx", make_docx())
    assert parsed.stats.get("tables") == 1
    text, chunks = chunk_document(parsed, source="handbook.docx")
    assert_offsets(text, chunks)

    tables = [c for c in chunks if c.kind == "table"]
    assert len(tables) > 1, "a large table should be split by rows"
    assert tables[0].text.startswith("| Site | Units | Interval |")
    assert all(t.table_header.startswith("| Site | Units | Interval |") for t in tables[1:]), "continuations carry the header"
    assert all(not t.text.startswith("| Site | Units") for t in tables[1:])

    rectifier = [c for c in chunks if c.kind == "text" and "Rectifier bank" in c.text]
    assert len(rectifier) > 1, "long prose splits on sentences"
    assert all(c.section == "Maintenance Handbook > Rectifiers" for c in rectifier)
    assert all(c.context.startswith("handbook.docx > Maintenance Handbook > Rectifiers") for c in rectifier)

    battery = [c for c in chunks if "batteries every four years" in c.text]
    assert battery and battery[0].section == "Maintenance Handbook > Batteries", "chunks never cross sections"
    assert "Short." in battery[0].text, "tiny trailing paragraphs merge into their neighbour"


def test_pdf_offsets_headings_pages_and_repeated_footer_removed():
    parsed = parse_file("pdf", make_pdf())
    assert parsed.page_count == 3
    headings = [e.text for e in parsed.elements if e.kind == "heading"]
    assert headings[:3] == ["Section 1 Overview", "Section 2 Overview", "Section 3 Overview"]
    text, chunks = chunk_document(parsed, source="manual.pdf")
    assert_offsets(text, chunks)
    assert {c.page for c in chunks} == {1, 2, 3}
    assert all(c.section.startswith("Section") for c in chunks)


def test_scanned_pdf_reports_pages_without_text():
    doc = pymupdf.open()
    doc.new_page()
    data = doc.tobytes()
    doc.close()
    parsed = parse_file("pdf", data)
    text, chunks = chunk_document(parsed, source="scan.pdf")
    assert parsed.stats.get("pages_without_text") == 1
    assert chunks == []


def test_unpunctuated_prose_is_split_on_words_with_exact_offsets():
    parsed = ParsedDocument(elements=[heading("Dump", 1), para(NO_PUNCTUATION)], doc_type="pdf")
    text, chunks = chunk_document(parsed, source="dump.pdf")
    assert_offsets(text, chunks)
    assert len(chunks) > 3
    assert max(estimate_tokens(c.text) for c in chunks) <= 820


def test_overlap_repeats_whole_sentences_within_a_section():
    parsed = ParsedDocument(elements=[heading("Intervals", 1), para(LONG_SENTENCE_PROSE)], doc_type="pdf")
    text, chunks = chunk_document(parsed, source="x.pdf")
    assert_offsets(text, chunks)
    for previous, current in zip(chunks, chunks[1:]):
        assert current.start < previous.end, "consecutive prose chunks overlap"
        overlap = text[current.start : previous.end]
        assert overlap.startswith("Rectifier bank"), "overlap starts at a sentence boundary"


def test_code_and_tables_are_never_merged_into_prose():
    parsed = ParsedDocument(
        elements=[
            heading("Setup", 1),
            para("Tiny."),
            code("def install():\n    return True"),
            para("Also tiny."),
            table("| a | b |\n| --- | --- |\n| 1 | 2 |"),
        ],
        doc_type="docx",
    )
    text, chunks = chunk_document(parsed, source="setup.docx")
    assert_offsets(text, chunks)
    kinds = [c.kind for c in chunks]
    assert "code" in kinds and "table" in kinds
    for chunk in chunks:
        if chunk.kind == "text":
            assert "def install" not in chunk.text and "| a | b |" not in chunk.text


def test_nul_bytes_are_removed_before_storage():
    parsed = ParsedDocument(elements=[para("Before\x00after " + LONG_SENTENCE_PROSE[:400])], doc_type="pdf")
    text, _ = chunk_document(parsed, source="nul.pdf")
    assert "\x00" not in text
